#!/usr/bin/env python3
"""Normalize mixed study sources into one JSONL stream with source anchors.

Supported without extra dependencies: TXT, MD, HTML, DOCX, DOC/RTF via a
local converter, EPUB, CSV, and TSV. XLSX is read from its XML parts. PDF and
PPTX are supported when PyMuPDF or python-pptx is installed.

Each output line describes one block: page, slide, paragraph, line range,
table row, or list item. Use the ref field in the coverage map.

Usage:
  python normalize_source.py sources/ --out working/normalized.jsonl
  python normalize_source.py deck.pdf notes.docx --out working/normalized.jsonl
  python normalize_source.py sources/ --out working/normalized.jsonl --ext .docx,.csv

Output:
  working/normalized.jsonl
  working/normalized.summary.json
"""
import argparse
import csv
import fnmatch
import html
import io
import json
import os
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote
import xml.etree.ElementTree as ET

TEXT_EXT = {".txt", ".text", ".md", ".markdown", ".mdown"}
HTML_EXT = {".html", ".htm", ".xhtml"}
DOCX_EXT = {".docx"}
LEGACY_DOC_EXT = {".doc", ".rtf"}
EPUB_EXT = {".epub"}
CSV_EXT = {".csv", ".tsv"}
XLSX_EXT = {".xlsx", ".xlsm"}
PDF_EXT = {".pdf"}
PPTX_EXT = {".pptx"}
SKIP_DIRS = {".git", "node_modules", "__pycache__", ".venv", "venv", "dist", "build", ".idea"}
BLOCK_TAGS = {
    "address", "article", "blockquote", "div", "dl", "dt", "dd", "figcaption",
    "figure", "h1", "h2", "h3", "h4", "h5", "h6", "li", "main", "ol", "p",
    "pre", "section", "table", "td", "th", "tr", "ul",
}
SKIP_TAGS = {
    "script", "style", "noscript", "template", "svg", "canvas", "iframe",
    "nav", "aside", "button", "form", "input", "select", "textarea",
    "footer",
}
VOID_TAGS = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}
QUIET = False


class AdapterError(Exception):
    pass


def log(message):
    if not QUIET:
        print(f"[normalize] {message}", file=sys.stderr)


def local_name(tag):
    return tag.rsplit("}", 1)[-1]


def attr(element, name):
    for key, value in element.attrib.items():
        if local_name(key) == name:
            return value
    return None


def safe_xml(data):
    head = data[:8192].lower()
    if b"<!doctype" in head or b"<!entity" in head:
        raise AdapterError("XML DTD or entity declarations are not allowed")
    return ET.fromstring(data)


MAX_ZIP_MEMBER = 50 * 1024 * 1024


def read_zip_member(archive, name):
    info = archive.getinfo(name)
    if info.file_size > MAX_ZIP_MEMBER:
        raise AdapterError(f"archive member too large: {name} ({info.file_size} bytes)")
    return archive.read(name)


def clean_text(value, keep_lines=False):
    value = html.unescape(value)
    value = value.replace("\u00a0", " ")
    if not keep_lines:
        value = re.sub(r"\s+", " ", value)
    else:
        value = re.sub(r"[ \t]+", " ", value)
        value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def display_path(path):
    resolved = path.resolve()
    try:
        relative = os.path.relpath(resolved, Path.cwd())
    except ValueError:
        return str(resolved)
    if relative.startswith(".."):
        return str(resolved)
    return relative


def read_text(path):
    raw = path.read_bytes()
    if b"\x00" in raw[:4096]:
        raise AdapterError("binary file with null bytes")
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise AdapterError("cannot decode text")


def split_long(text, limit):
    text = text.strip()
    if len(text) <= limit:
        return [text]
    parts = []
    rest = text
    while len(rest) > limit:
        window = rest[:limit]
        cut = max(window.rfind("\n"), window.rfind(". "), window.rfind("; "), window.rfind(", "))
        if cut < limit // 3:
            cut = window.rfind(" ")
        if cut < limit // 3:
            cut = limit
        parts.append(rest[:cut].strip())
        rest = rest[cut:].strip()
    if rest:
        parts.append(rest)
    return [part for part in parts if part]


def make_record(source, fmt, unit, ref, text, section=None, meta=None, keep_lines=False):
    record = {
        "source": source,
        "format": fmt,
        "unit": unit,
        "ref": ref,
        "section": section,
        "text": clean_text(text, keep_lines=keep_lines),
    }
    if meta:
        record["meta"] = meta
    return record


def add_split(records, source, fmt, unit, base_ref, text, section=None, meta=None, limit=4000, keep_lines=False):
    parts = split_long(text, limit)
    for index, part in enumerate(parts, 1):
        ref = base_ref if len(parts) == 1 else f"{base_ref}!part:{index}"
        extra = dict(meta or {})
        if len(parts) > 1:
            extra["part"] = index
            extra["parts"] = len(parts)
        records.append(make_record(source, fmt, unit, ref, part, section, extra or None, keep_lines))


def line_ref(start, end=None):
    if end and end != start:
        return f"line:{start}-{end}"
    return f"line:{start}"


def column_letter(index):
    letters = ""
    value = index
    while value > 0:
        value, remainder = divmod(value - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters or "A"


def update_sections(stack, level, title):
    stack = stack[: max(level - 1, 0)]
    while len(stack) < level - 1:
        stack.append("")
    stack.append(title)
    return stack


def parse_plain_text(text, source, fmt, limit):
    records = []
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    buffer = []
    start = None
    end = None

    def flush():
        nonlocal buffer, start, end
        if buffer:
            add_split(records, source, fmt, "paragraph", line_ref(start, end), " ".join(buffer), None, None, limit)
        buffer = []
        start = None
        end = None

    for number, line in enumerate(lines, 1):
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if start is None:
            start = number
        end = number
        buffer.append(stripped)
    flush()
    log(f"{fmt}: {display_path(Path(source))} -> {len(records)} paragraphs")
    return records, "plain text"


def strip_markdown_inline(value):
    value = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", value)
    value = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", value)
    value = re.sub(r"`([^`]+)`", r"\1", value)
    value = re.sub(r"\*\*(.+?)\*\*", r"\1", value)
    value = re.sub(r"__(.+?)__", r"\1", value)
    value = re.sub(r"~~(.+?)~~", r"\1", value)
    value = re.sub(r"(?<!\w)\*([^*]+)\*(?!\w)", r"\1", value)
    value = re.sub(r"(?<!\w)_([^_]+)_(?!\w)", r"\1", value)
    value = re.sub(r"<[^>]+>", " ", value)
    return clean_text(value)


def split_markdown_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [strip_markdown_inline(cell) for cell in line.split("|")]


def is_table_separator(line):
    cells = split_markdown_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", cell.replace(" ", "")) for cell in cells)


def parse_markdown(text, source, limit):
    records = []
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    index = 0
    if lines and lines[0].strip() == "---":
        for probe in range(1, min(len(lines), 200)):
            if lines[probe].strip() == "---":
                index = probe + 1
                break
    section_stack = []
    buffer = []
    start = None
    end = None
    in_code = False
    code_lines = []
    code_start = None
    code_language = ""

    def section():
        return " > ".join(part for part in section_stack if part) or None

    def flush(last_line):
        nonlocal buffer, start, end
        if buffer:
            add_split(records, source, "md", "paragraph", line_ref(start, end or start), " ".join(buffer), section(), None, limit)
        buffer = []
        start = None
        end = None

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        number = index + 1

        if stripped.startswith("```"):
            flush(number - 1)
            if in_code:
                add_split(records, source, "md", "code", line_ref(code_start, number), "\n".join(code_lines), section(), {"language": code_language}, limit, True)
                in_code = False
                code_lines = []
                code_start = None
                code_language = ""
            else:
                in_code = True
                code_start = number
                code_language = stripped[3:].strip()
            index += 1
            continue

        if in_code:
            code_lines.append(raw.rstrip())
            index += 1
            continue

        if not stripped:
            flush(number - 1)
            index += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            flush(number - 1)
            level = len(heading.group(1))
            title = strip_markdown_inline(heading.group(2))
            section_stack = update_sections(section_stack, level, title)
            current = section()
            records.append(make_record(source, "md", "heading", line_ref(number), title, current))
            index += 1
            continue

        if "|" in stripped and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            flush(number - 1)
            table_rows = []
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                row_number = index + 1
                if not is_table_separator(lines[index]):
                    table_rows.append((row_number, split_markdown_row(lines[index])))
                index += 1
            if table_rows:
                headers = table_rows[0][1]
                header_number = table_rows[0][0]
                records.append(
                    make_record(
                        source,
                        "md",
                        "header",
                        line_ref(header_number),
                        " | ".join(headers),
                        section(),
                        {"columns": headers, "header": True},
                    )
                )
                for row_number, cells in table_rows[1:]:
                    width = max(len(headers), len(cells))
                    padded_headers = headers + [f"Column {column_letter(i)}" for i in range(len(headers) + 1, width + 1)]
                    padded = cells + [""] * (width - len(cells))
                    text_row = " | ".join(
                        f"{padded_headers[i]}: {value}" if padded_headers[i] and value else (value or padded_headers[i])
                        for i, value in enumerate(padded)
                        if value or padded_headers[i]
                    )
                    records.append(make_record(source, "md", "row", line_ref(row_number), text_row, section(), {"cells": padded, "columns": padded_headers}))
            continue

        item = re.match(r"^([-*+]|\d+[.)])\s+(.*)$", stripped)
        if item:
            flush(number - 1)
            records.append(make_record(source, "md", "list_item", line_ref(number), strip_markdown_inline(item.group(2)), section()))
            index += 1
            continue

        if stripped.startswith(">"):
            stripped = stripped.lstrip(">").strip()
            if not stripped:
                index += 1
                continue

        if start is None:
            start = number
        end = number
        buffer.append(strip_markdown_inline(stripped))
        index += 1

    flush(len(lines))
    if in_code and code_lines:
        add_split(records, source, "md", "code", line_ref(code_start, len(lines)), "\n".join(code_lines), section(), {"language": code_language}, limit, True)
    log(f"md: {display_path(Path(source))} -> {len(records)} blocks")
    return records, "markdown"


class HtmlExtractor(HTMLParser):
    def __init__(self, source, ref_prefix="", initial_section=None, limit=4000, fmt="html"):
        super().__init__(convert_charrefs=True)
        self.source = source
        self.fmt = fmt
        self.ref_prefix = ref_prefix
        self.section_stack = [initial_section] if initial_section else []
        self.limit = limit
        self.records = []
        self.buffer = []
        self.block_start = None
        self.block_end = None
        self.skip_stack = []
        self.heading_level = None
        self.heading_text = []
        self.heading_start = None
        self.in_title = False
        self.title_text = []
        self.in_cell = False
        self.cell_text = []
        self.current_row = []
        self.current_table = []
        self.table_counter = 0
        self.li_depth = 0

    def section(self):
        return " > ".join(part for part in self.section_stack if part) or None

    def position(self):
        return self.getpos()[0]

    def ref(self, start, end=None):
        base = line_ref(start, end)
        return f"{self.ref_prefix}{base}" if self.ref_prefix else base

    def flush(self, unit="paragraph", last_line=None):
        if not self.buffer:
            return
        text = clean_text(" ".join(self.buffer))
        start = self.block_start or 1
        end = last_line or self.block_end or start
        if text:
            add_split(self.records, self.source, self.fmt, unit, self.ref(start, end), text, self.section(), None, self.limit)
        self.buffer = []
        self.block_start = None
        self.block_end = None

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if self.skip_stack:
            if tag in SKIP_TAGS and tag not in VOID_TAGS:
                self.skip_stack.append(tag)
            return
        if tag in SKIP_TAGS:
            if tag not in VOID_TAGS:
                self.skip_stack.append(tag)
            return
        if self.in_cell:
            if tag == "br":
                self.cell_text.append(" ")
            return
        if tag in {"td", "th"}:
            self.in_cell = True
            self.cell_text = []
            return
        if tag == "tr":
            self.current_row = []
            return
        if tag == "table":
            self.flush("paragraph", self.position() - 1)
            self.table_counter += 1
            self.current_table = []
            return
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self.flush("paragraph", self.position() - 1)
            self.heading_level = int(tag[1])
            self.heading_text = []
            self.heading_start = self.position()
            return
        if tag == "title":
            self.in_title = True
            self.title_text = []
            return
        if tag == "li":
            self.flush("paragraph", self.position() - 1)
            self.li_depth += 1
            return
        if tag == "br":
            if self.heading_level:
                self.heading_text.append(" ")
            else:
                self.buffer.append(" ")
            return
        if tag in BLOCK_TAGS:
            self.flush("paragraph", self.position() - 1)

    def handle_endtag(self, tag):
        tag = tag.lower()
        if self.skip_stack:
            if tag == self.skip_stack[-1]:
                self.skip_stack.pop()
            return
        if tag in {"td", "th"} and self.in_cell:
            self.current_row.append(clean_text(" ".join(self.cell_text)))
            self.in_cell = False
            self.cell_text = []
            return
        if tag == "tr" and self.current_row:
            self.current_table.append((self.position(), self.current_row))
            self.current_row = []
            return
        if tag == "table":
            self.emit_table()
            return
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"} and self.heading_level:
            title = clean_text(" ".join(self.heading_text))
            if title:
                self.section_stack = update_sections(self.section_stack, self.heading_level, title)
                self.records.append(make_record(self.source, self.fmt, "heading", self.ref(self.heading_start), title, self.section()))
            self.heading_level = None
            self.heading_text = []
            return
        if tag == "title":
            self.in_title = False
            return
        if tag == "li" and self.li_depth:
            self.li_depth -= 1
            self.flush("list_item", self.position())
            return
        if tag in BLOCK_TAGS:
            unit = "list_item" if self.li_depth and tag == "p" else "paragraph"
            self.flush(unit, self.position())

    def handle_data(self, data):
        if self.skip_stack:
            return
        if self.in_title:
            self.title_text.append(data)
            return
        if self.in_cell:
            self.cell_text.append(data)
            return
        if self.heading_level:
            self.heading_text.append(data)
            return
        if not data.strip():
            return
        if self.block_start is None:
            self.block_start = self.position()
        self.block_end = self.position()
        self.buffer.append(data)

    def emit_table(self):
        if not self.current_table:
            return
        header_number, header = self.current_table[0]
        width = max(len(cells) for _, cells in self.current_table)
        headers = header + [f"Column {column_letter(i)}" for i in range(len(header) + 1, width + 1)]
        meta = {"columns": headers, "table": self.table_counter, "header": True}
        self.records.append(make_record(self.source, self.fmt, "header", self.ref(header_number), " | ".join(headers), self.section(), meta))
        for row_number, cells in self.current_table[1:]:
            padded = cells + [""] * (width - len(cells))
            text_row = " | ".join(
                f"{headers[index]}: {value}" if headers[index] and value else (value or headers[index])
                for index, value in enumerate(padded)
                if value or headers[index]
            )
            if not text_row:
                continue
            meta = {"cells": padded, "columns": headers, "table": self.table_counter}
            self.records.append(make_record(self.source, self.fmt, "row", self.ref(row_number), text_row, self.section(), meta))
        self.current_table = []

    def finish(self):
        self.flush("paragraph", self.position())
        self.emit_table()
        return clean_text(" ".join(self.title_text)) or None


def parse_html(text, source, ref_prefix="", initial_section=None, limit=4000, fmt="html"):
    parser = HtmlExtractor(source, ref_prefix, initial_section, limit, fmt)
    parser.feed(text)
    parser.close()
    title = parser.finish()
    log(f"html: {source} -> {len(parser.records)} blocks")
    return parser.records, title


def parse_epub(path, limit):
    source = display_path(path)
    records = []
    chapters = []
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        opf_path = None
        if "META-INF/container.xml" in names:
            container = safe_xml(read_zip_member(archive, "META-INF/container.xml"))
            for element in container.iter():
                if local_name(element.tag) == "rootfile":
                    opf_path = attr(element, "full-path")
                    break
        if not opf_path:
            raise AdapterError("EPUB has no container rootfile")
        base = posixpath.dirname(opf_path)
        package = safe_xml(read_zip_member(archive, opf_path))
        manifest = {}
        for element in package.iter():
            if local_name(element.tag) == "item":
                manifest[attr(element, "id")] = (attr(element, "href"), attr(element, "media-type"))
        spine = []
        for element in package.iter():
            if local_name(element.tag) == "itemref":
                spine.append(attr(element, "idref"))
        ordered = []
        for idref in spine:
            item = manifest.get(idref)
            if item and item[1] in {"application/xhtml+xml", "text/html"}:
                ordered.append(item[0])
        if not ordered:
            ordered = sorted(href for href, media in manifest.values() if media and "html" in media)
        if not ordered:
            ordered = sorted(name for name in names if name.lower().endswith((".xhtml", ".html")))
        for index, href in enumerate(ordered, 1):
            chapter_path = posixpath.normpath(posixpath.join(base, unquote(href.split("#")[0])))
            if chapter_path not in names:
                log(f"epub: missing {chapter_path}")
                continue
            chapter_html = read_zip_member(archive, chapter_path).decode("utf-8", errors="replace")
            title_match = re.search(r"<title[^>]*>(.*?)</title>", chapter_html, re.I | re.S)
            chapter_title = clean_text(html.unescape(title_match.group(1))) if title_match else f"Chapter {index}"
            chapter_records, _ = parse_html(
                chapter_html,
                source,
                ref_prefix=f"chapter:{index}!",
                initial_section=chapter_title,
                limit=limit,
                fmt="epub",
            )
            records.extend(chapter_records)
            chapters.append({"index": index, "title": chapter_title, "href": href, "records": len(chapter_records)})
    log(f"epub: {source} -> {len(chapters)} chapters, {len(records)} blocks")
    return records, f"{len(chapters)} chapters"


def word_text(paragraph):
    parts = []
    for element in paragraph.iter():
        name = local_name(element.tag)
        if name == "t" and element.text:
            parts.append(element.text)
        elif name == "tab":
            parts.append("\t")
        elif name in {"br", "cr"}:
            parts.append("\n")
    return clean_text("".join(parts), keep_lines=True)


def word_style(paragraph):
    for child in paragraph:
        if local_name(child.tag) == "pPr":
            for item in child:
                if local_name(item.tag) == "pStyle":
                    return attr(item, "val") or ""
    return ""


def heading_level(style):
    if not style:
        return 0
    match = re.match(r"(?i)heading\s*([1-6])", style)
    if match:
        return int(match.group(1))
    if style.lower() in {"title", "subtitle"}:
        return 1
    return 0


def iter_word_blocks(parent):
    for child in parent:
        name = local_name(child.tag)
        if name in {"p", "tbl"}:
            yield name, child
        elif name in {"sdt", "sdtContent", "customXml"}:
            yield from iter_word_blocks(child)


def table_cells(row):
    cells = []
    for cell in row:
        if local_name(cell.tag) != "tc":
            continue
        texts = [word_text(paragraph) for paragraph in cell if local_name(paragraph.tag) == "p"]
        cells.append(clean_text(" ".join(text for text in texts if text)))
    return cells


def parse_docx(path, limit):
    source = display_path(path)
    with zipfile.ZipFile(path) as archive:
        if "word/document.xml" not in archive.namelist():
            raise AdapterError("DOCX has no word/document.xml")
        root = safe_xml(read_zip_member(archive, "word/document.xml"))
    body = None
    for element in root.iter():
        if local_name(element.tag) == "body":
            body = element
            break
    if body is None:
        raise AdapterError("DOCX body not found")
    records = []
    section_stack = []
    paragraph_number = 0
    table_number = 0
    for kind, block in iter_word_blocks(body):
        if kind == "p":
            paragraph_number += 1
            text = word_text(block)
            if not text:
                continue
            style = word_style(block)
            level = heading_level(style)
            if level:
                section_stack = update_sections(section_stack, level, text)
            current_section = " > ".join(part for part in section_stack if part) or None
            unit = "heading" if level else "paragraph"
            meta = {"style": style} if style else None
            add_split(records, source, "docx", unit, f"paragraph:{paragraph_number}", text, current_section, meta, limit)
        else:
            table_number += 1
            rows = [row for row in block if local_name(row.tag) == "tr"]
            width = 0
            parsed_rows = []
            for row in rows:
                cells = table_cells(row)
                width = max(width, len(cells))
                parsed_rows.append(cells)
            headers = parsed_rows[0] if parsed_rows else []
            headers = headers + [f"Column {column_letter(i)}" for i in range(len(headers) + 1, width + 1)]
            current_section = " > ".join(part for part in section_stack if part) or None
            if headers:
                add_split(
                    records,
                    source,
                    "docx",
                    "header",
                    f"table:{table_number}!row:1",
                    " | ".join(headers),
                    current_section,
                    {"columns": headers, "table": table_number, "header": True},
                    limit,
                )
            for row_index, cells in enumerate(parsed_rows[1:], 2):
                padded = cells + [""] * (width - len(cells))
                text_row = " | ".join(
                    f"{headers[index]}: {value}" if headers[index] and value else (value or headers[index])
                    for index, value in enumerate(padded)
                    if value or headers[index]
                )
                if not text_row:
                    continue
                add_split(
                    records,
                    source,
                    "docx",
                    "row",
                    f"table:{table_number}!row:{row_index}",
                    text_row,
                    current_section,
                    {"cells": padded, "columns": headers, "table": table_number},
                    limit,
                )
    if not records:
        raise AdapterError("DOCX produced no text")
    log(f"docx: {source} -> {len(records)} blocks, {table_number} tables")
    return records, f"{paragraph_number} paragraphs, {table_number} tables"


def convert_legacy(path):
    textutil = shutil.which("textutil")
    if textutil:
        try:
            result = subprocess.run(
                [textutil, "-convert", "txt", "-stdout", str(path)],
                capture_output=True,
                text=True,
                timeout=180,
                check=False,
            )
            if result.returncode == 0 and result.stdout.strip():
                return result.stdout, "textutil"
            if result.stderr.strip():
                log(f"textutil: {result.stderr.strip().splitlines()[0]}")
        except (OSError, subprocess.SubprocessError) as exc:
            log(f"textutil failed: {exc}")
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        with tempfile.TemporaryDirectory() as directory:
            try:
                result = subprocess.run(
                    [soffice, "--headless", "--convert-to", "txt:Text", "--outdir", directory, str(path)],
                    capture_output=True,
                    text=True,
                    timeout=240,
                    check=False,
                )
            except (OSError, subprocess.SubprocessError) as exc:
                raise AdapterError(f"LibreOffice failed: {exc}") from exc
            output = Path(directory) / f"{path.stem}.txt"
            if result.returncode == 0 and output.exists():
                return read_text(output), "libreoffice"
            raise AdapterError("LibreOffice conversion produced no text")
    raise AdapterError("no DOC converter; install LibreOffice or convert to DOCX first")


def parse_legacy_doc(path, limit):
    source = display_path(path)
    text, converter = convert_legacy(path)
    records, note = parse_plain_text(text, source, "doc", limit)
    log(f"doc: {source} converted via {converter}")
    return records, f"converted via {converter}; {note}"


def table_rows_from_matrix(rows, source, fmt, sheet_name, limit, ref_prefix=""):
    records = []
    numbered = []
    for index, item in enumerate(rows, 1):
        if isinstance(item, tuple):
            number, values = item
        else:
            number, values = index, item
        numbered.append((number, values))
    if not numbered:
        return records
    header_index = next((index for index, (_, row) in enumerate(numbered) if any(value.strip() for value in row)), None)
    if header_index is None:
        return records
    header_number, header_row = numbered[header_index]
    width = max(len(row) for _, row in numbered)
    headers = []
    seen = {}
    for index in range(width):
        label = clean_text(header_row[index]) if index < len(header_row) else ""
        if not label:
            label = f"Column {column_letter(index + 1)}"
        base = label
        seen[base] = seen.get(base, 0) + 1
        if seen[base] > 1:
            label = f"{base} ({seen[base]})"
        headers.append(label)
    header_meta = {"columns": headers, "header": True}
    if sheet_name:
        header_meta["sheet"] = sheet_name
    add_split(
        records,
        source,
        fmt,
        "header",
        f"{ref_prefix}row:{header_number}",
        " | ".join(headers),
        sheet_name,
        header_meta,
        limit,
    )
    for number, row in numbered[header_index + 1:]:
        padded = [clean_text(value) for value in row] + [""] * (width - len(row))
        if not any(padded):
            continue
        text_row = " | ".join(
            f"{headers[index]}: {value}" if headers[index] and value else (value or headers[index])
            for index, value in enumerate(padded)
            if value or headers[index]
        )
        if not text_row:
            continue
        meta = {"cells": padded, "columns": headers, "row": number}
        if sheet_name:
            meta["sheet"] = sheet_name
        add_split(
            records,
            source,
            fmt,
            "row",
            f"{ref_prefix}row:{number}",
            text_row,
            sheet_name,
            meta,
            limit,
        )
    return records


def parse_csv(path, limit):
    source = display_path(path)
    text = read_text(path)
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        pass
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = []
    for row in reader:
        rows.append([value.strip() for value in row])
    records = table_rows_from_matrix(rows, source, "csv", None, limit)
    log(f"csv: {source} -> {len(rows)} rows, delimiter={repr(delimiter)}")
    return records, f"{len(rows)} rows, delimiter {repr(delimiter)}"


def xlsx_shared_strings(archive):
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []
    root = safe_xml(read_zip_member(archive, "xl/sharedStrings.xml"))
    values = []
    for item in root:
        if local_name(item.tag) != "si":
            continue
        texts = [element.text or "" for element in item.iter() if local_name(element.tag) == "t"]
        values.append("".join(texts))
    return values


def xlsx_sheet_targets(archive):
    workbook = safe_xml(read_zip_member(archive, "xl/workbook.xml"))
    relationships = {}
    if "xl/_rels/workbook.xml.rels" in archive.namelist():
        rel_root = safe_xml(read_zip_member(archive, "xl/_rels/workbook.xml.rels"))
        for relationship in rel_root:
            relationships[attr(relationship, "Id")] = attr(relationship, "Target")
    sheets = []
    for element in workbook.iter():
        if local_name(element.tag) != "sheet":
            continue
        name = attr(element, "name") or "Sheet"
        rel_id = attr(element, "id")
        target = relationships.get(rel_id, "")
        if target.startswith("/"):
            target = target.lstrip("/")
        else:
            target = posixpath.normpath(posixpath.join("xl", target))
        sheets.append((name, target))
    return sheets


def xlsx_cell_value(cell, shared):
    cell_type = attr(cell, "t") or ""
    value = None
    inline = []
    for child in cell:
        name = local_name(child.tag)
        if name == "v":
            value = child.text
        elif name == "is":
            inline = [element.text or "" for element in child.iter() if local_name(element.tag) == "t"]
    if cell_type == "s" and value is not None:
        try:
            return shared[int(value)]
        except (ValueError, IndexError):
            return ""
    if cell_type == "inlineStr":
        return "".join(inline)
    if cell_type == "b":
        return "TRUE" if value == "1" else "FALSE"
    if cell_type == "e":
        return value or ""
    return value or ""


def parse_xlsx(path, limit):
    source = display_path(path)
    records = []
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if "xl/workbook.xml" not in names:
            raise AdapterError("XLSX has no workbook.xml")
        shared = xlsx_shared_strings(archive)
        for sheet_name, target in xlsx_sheet_targets(archive):
            if target not in names:
                log(f"xlsx: missing sheet target {target}")
                continue
            root = safe_xml(read_zip_member(archive, target))
            rows = []
            for row in root.iter():
                if local_name(row.tag) != "row":
                    continue
                row_number = int(attr(row, "r") or len(rows) + 1)
                values = []
                for cell in row:
                    if local_name(cell.tag) != "c":
                        continue
                    reference = attr(cell, "r") or ""
                    letters = re.match(r"([A-Z]+)", reference.upper())
                    column = 0
                    if letters:
                        for character in letters.group(1):
                            column = column * 26 + (ord(character) - 64)
                    else:
                        column = len(values) + 1
                    while len(values) < column - 1:
                        values.append("")
                    values.append(xlsx_cell_value(cell, shared))
                while values and values[-1] == "":
                    values.pop()
                rows.append((row_number, values))
            sheet_records = table_rows_from_matrix(rows, source, "xlsx", sheet_name, limit, ref_prefix=f"{sheet_name}!")
            records.extend(sheet_records)
            log(f"xlsx: {sheet_name} -> {len(sheet_records)} rows")
    if not records:
        raise AdapterError("XLSX produced no rows")
    log(f"xlsx: {source} -> {len(records)} rows total")
    return records, f"{len(records)} rows"


def parse_pdf(path, limit):
    try:
        import pymupdf
    except ImportError as exc:
        raise AdapterError("PyMuPDF not installed; pip install pymupdf") from exc
    source = display_path(path)
    records = []
    empty = 0
    document = pymupdf.open(str(path))
    try:
        for number, page in enumerate(document, 1):
            text = page.get_text().strip()
            if not text:
                empty += 1
                continue
            add_split(records, source, "pdf", "page", f"page:{number}", text, None, None, limit)
    finally:
        document.close()
    if not records:
        raise AdapterError("PDF has no text; run OCR first")
    log(f"pdf: {source} -> {len(records)} pages, empty={empty}")
    return records, f"{len(records)} pages, {empty} empty"


def parse_pptx(path, limit):
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise AdapterError("python-pptx not installed; pip install python-pptx") from exc
    source = display_path(path)
    records = []
    presentation = Presentation(str(path))
    for number, slide in enumerate(presentation.slides, 1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        text = clean_text("\n".join(texts), keep_lines=True)
        if text:
            add_split(records, source, "pptx", "slide", f"slide:{number}", text, None, None, limit)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                add_split(records, source, "pptx", "note", f"slide:{number}!notes", notes, None, None, limit)
    if not records:
        raise AdapterError("PPTX has no text")
    log(f"pptx: {source} -> {len(records)} blocks")
    return records, f"{len(records)} blocks"


def unsupported_hint(suffix):
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}:
        return "image; run scripts/ocr.sh"
    if suffix in {".m4a", ".mp3", ".wav", ".aac", ".mp4", ".mov", ".mkv"}:
        return "audio or video; transcribe with Whisper, then normalize the transcript"
    if suffix in {".ppt", ".xls"}:
        return "legacy Office format; convert to PPTX or XLSX first"
    if suffix == ".srt" or suffix == ".vtt":
        return "transcript; convert with scripts/srt_to_paragraphs.py, then normalize the text"
    return "unsupported format"


def parse_file(path, limit):
    suffix = path.suffix.lower()
    if suffix in TEXT_EXT:
        text = read_text(path)
        if suffix in {".md", ".markdown", ".mdown"}:
            return parse_markdown(text, display_path(path), limit)
        return parse_plain_text(text, display_path(path), "txt", limit)
    if suffix in HTML_EXT:
        return parse_html(read_text(path), display_path(path), limit=limit)
    if suffix in DOCX_EXT:
        return parse_docx(path, limit)
    if suffix in LEGACY_DOC_EXT:
        return parse_legacy_doc(path, limit)
    if suffix in EPUB_EXT:
        return parse_epub(path, limit)
    if suffix in CSV_EXT:
        return parse_csv(path, limit)
    if suffix in XLSX_EXT:
        return parse_xlsx(path, limit)
    if suffix in PDF_EXT:
        return parse_pdf(path, limit)
    if suffix in PPTX_EXT:
        return parse_pptx(path, limit)
    raise AdapterError(unsupported_hint(suffix))


def collect_files(paths, extensions, excludes):
    files = []
    skipped = []
    seen = set()
    for raw in paths:
        path = Path(raw).expanduser()
        if not path.exists():
            skipped.append({"path": str(path), "reason": "not found"})
            continue
        candidates = [path] if path.is_file() else sorted(path.rglob("*"))
        for candidate in candidates:
            if not candidate.is_file():
                continue
            if candidate.name.startswith("~$") or candidate.name.startswith("."):
                continue
            if any(part in SKIP_DIRS for part in candidate.parts):
                continue
            display = display_path(candidate)
            if any(fnmatch.fnmatch(display, pattern) or fnmatch.fnmatch(candidate.name, pattern) for pattern in excludes):
                skipped.append({"path": display, "reason": "excluded"})
                continue
            suffix = candidate.suffix.lower()
            if extensions and suffix not in extensions:
                skipped.append({"path": display, "reason": f"extension {suffix or '(none)'} filtered out"})
                continue
            key = candidate.resolve()
            if key in seen:
                continue
            seen.add(key)
            files.append(candidate)
    return files, skipped


def write_outputs(records, summary, out_path, summary_path):
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")
    if summary_path:
        summary_path.parent.mkdir(parents=True, exist_ok=True)
        summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    global QUIET
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("paths", nargs="+", help="Files or directories to normalize")
    parser.add_argument("--out", type=Path, required=True, help="JSONL output path")
    parser.add_argument("--summary", type=Path, default=None, help="Summary JSON path; defaults to <out>.summary.json")
    parser.add_argument("--no-summary", action="store_true", help="Do not write a summary file")
    parser.add_argument("--ext", default=None, help="Comma-separated extension allowlist, for example .docx,.csv")
    parser.add_argument("--exclude", action="append", default=[], help="Glob to skip; repeatable")
    parser.add_argument("--max-chars", type=int, default=4000, help="Split blocks longer than this")
    parser.add_argument("--quiet", action="store_true", help="Suppress debug logs")
    args = parser.parse_args()
    QUIET = args.quiet

    extensions = None
    if args.ext:
        extensions = {item.strip().lower() if item.strip().startswith(".") else f".{item.strip().lower()}" for item in args.ext.split(",") if item.strip()}

    files, skipped = collect_files(args.paths, extensions, args.exclude)
    log(f"scan: {len(files)} files, {len(skipped)} skipped, max_chars={args.max_chars}")

    records = []
    sources = []
    errors = 0
    converters = set()
    for file in files:
        suffix = file.suffix.lower()
        supported = suffix in TEXT_EXT | HTML_EXT | DOCX_EXT | LEGACY_DOC_EXT | EPUB_EXT | CSV_EXT | XLSX_EXT | PDF_EXT | PPTX_EXT
        display = display_path(file)
        if not supported:
            skipped.append({"path": display, "reason": unsupported_hint(suffix)})
            log(f"skip: {display} ({unsupported_hint(suffix)})")
            continue
        try:
            file_records, note = parse_file(file, args.max_chars)
            file_chars = sum(len(record["text"]) for record in file_records)
            records.extend(file_records)
            if suffix in LEGACY_DOC_EXT:
                converters.add(note.split("via ", 1)[-1].split(";", 1)[0] if "via " in note else "unknown")
            sources.append(
                {
                    "source": display,
                    "format": suffix.lstrip("."),
                    "status": "ok",
                    "units": len(file_records),
                    "chars": file_chars,
                    "note": note,
                }
            )
            log(f"ok: {display} -> {len(file_records)} records, {file_chars} chars")
        except (AdapterError, OSError, ValueError, zipfile.BadZipFile, ET.ParseError) as exc:
            errors += 1
            sources.append(
                {
                    "source": display,
                    "format": suffix.lstrip("."),
                    "status": "error",
                    "units": 0,
                    "chars": 0,
                    "note": str(exc),
                }
            )
            log(f"error: {display}: {exc}")

    unit_counts = {}
    format_counts = {}
    for record in records:
        unit_counts[record["unit"]] = unit_counts.get(record["unit"], 0) + 1
        format_counts[record["format"]] = format_counts.get(record["format"], 0) + 1

    summary = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "out": str(args.out),
        "max_chars": args.max_chars,
        "sources": sources,
        "totals": {
            "files": len(sources),
            "records": len(records),
            "chars": sum(len(record["text"]) for record in records),
            "errors": errors,
        },
        "units": unit_counts,
        "formats": format_counts,
        "converters": sorted(converters),
        "skipped": skipped[:200],
        "skipped_total": len(skipped),
    }
    summary_path = None if args.no_summary else (args.summary or args.out.with_suffix(".summary.json"))
    write_outputs(records, summary, args.out, summary_path)

    print(f"files={len(sources)} records={len(records)} chars={summary['totals']['chars']} errors={errors}")
    print(f"units={unit_counts}")
    print(f"wrote {args.out}")
    if summary_path:
        print(f"wrote {summary_path}")
    if not records:
        print("No usable records produced.", file=sys.stderr)
        return 1
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
